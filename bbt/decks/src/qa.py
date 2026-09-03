import sys, math
from pptx import Presentation
from pptx.util import Emu

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5
# avg glyph width as a fraction of point size (Calibri/Cambria are close)
GW = {"Cambria": 0.50, "Calibri": 0.47}

def inches(v): return v / 914400.0 if v is not None else None

def text_blocks(sh):
    if not sh.has_text_frame: return None
    tf = sh.text_frame
    runs = [(r.text, (r.font.size.pt if r.font.size else 18), (r.font.name or "Calibri"))
            for p in tf.paragraphs for r in p.runs]
    if not runs: return None
    npara = len([p for p in tf.paragraphs if p.runs])
    return runs, npara

def est_height(runs, npara, w_in):
    """rough rendered height in inches"""
    total = 0.0
    # group runs per paragraph approximation: treat all runs as one flow per paragraph
    per = max(1, len(runs) // max(1, npara))
    chunks = [runs[i:i+per] for i in range(0, len(runs), per)] or [runs]
    for ch in chunks:
        txt = "".join(t for t, _, _ in ch)
        size = max(s for _, s, _ in ch)
        face = ch[0][2]
        gw = GW.get(face, 0.48) * size
        cpl = max(4, int((w_in * 72) / gw))
        # honour explicit newlines
        lines = 0
        for seg in txt.split("\n"):
            lines += max(1, math.ceil(len(seg) / cpl))
        total += lines * size * 1.22 / 72.0
    return total

def rects_overlap(a, b, tol=0.02):
    ax, ay, aw, ah = a; bx, by, bw, bh = b
    ix = min(ax+aw, bx+bw) - max(ax, bx)
    iy = min(ay+ah, by+bh) - max(ay, by)
    return ix > tol and iy > tol, ix, iy

def main(path):
    prs = Presentation(path)
    issues = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in slide.shapes:
            x, y = inches(sh.left), inches(sh.top)
            w, h = inches(sh.width), inches(sh.height)
            if None in (x, y, w, h): continue
            name = (sh.shape_type or "")
            # bounds
            if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
                issues.append(f"s{i}: OUT OF BOUNDS  x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}")
            tb = text_blocks(sh)
            if tb:
                runs, npara = tb
                need = est_height(runs, npara, w)
                preview = "".join(t for t,_,_ in runs)[:44].replace("\n", " ")
                if need > h + 0.06:
                    issues.append(f"s{i}: OVERFLOW  need {need:.2f}\" have {h:.2f}\"  [{preview}]")
                if y + need > SLIDE_H - 0.15:
                    issues.append(f"s{i}: RUNS OFF BOTTOM  ends {y+need:.2f}\"  [{preview}]")
                texts.append(((x, y, w, min(h, need) if need < h else need), preview))
        # text-vs-text overlap
        for a in range(len(texts)):
            for b in range(a+1, len(texts)):
                ov, ix, iy = rects_overlap(texts[a][0], texts[b][0], tol=0.05)
                if ov and ix > 0.25 and iy > 0.12:
                    issues.append(f"s{i}: TEXT OVERLAP {ix:.2f}x{iy:.2f}  [{texts[a][1][:26]}] / [{texts[b][1][:26]}]")
    print(f"{path}: {len(prs.slides.__iter__.__self__._sldIdLst)} slides, {len(issues)} issues")
    for x in issues[:40]: print("  " + x)
    return len(issues)

if __name__ == "__main__":
    tot = 0
    for p in sys.argv[1:]: tot += main(p)
    sys.exit(0)
