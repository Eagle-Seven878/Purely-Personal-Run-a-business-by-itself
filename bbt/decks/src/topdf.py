"""Render a .pptx to HTML using real shape geometry, then to PDF via Chromium."""
import sys, os, html, pathlib
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400.0
def I(v): return (v or 0) / EMU

def rgb(c):
    try:
        if c and c.type is not None and c.rgb is not None: return "#%s" % str(c.rgb)
    except Exception: pass
    return None

def shape_fill(sh):
    try:
        f = sh.fill
        if f.type is not None and f.type == 1:  # solid
            return rgb(f.fore_color)
    except Exception: pass
    return None

def para_html(p):
    align = {1: "center", 2: "right", 3: "justify"}.get(p.alignment, "left")
    bullet = False
    try:
        ppr = p._pPr
        if ppr is not None and ppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') is not None:
            bullet = True
    except Exception: pass
    runs = []
    for r in p.runs:
        f = r.font
        st = []
        if f.size: st.append("font-size:%.2fpt" % f.size.pt)
        if f.bold: st.append("font-weight:700")
        if f.italic: st.append("font-style:italic")
        c = rgb(f.color)
        if c: st.append("color:%s" % c)
        if f.name:
            fam = ("Cambria,'DejaVu Serif',Georgia,serif" if 'Cambria' in f.name
                   else "Calibri,'DejaVu Sans',Arial,sans-serif")
            st.append("font-family:%s" % fam)
        runs.append('<span style="%s">%s</span>' % (";".join(st), html.escape(r.text)))
    if not runs: return ""
    ls = "line-height:1.18;"
    try:
        if p.line_spacing and not isinstance(p.line_spacing, float):
            ls = "line-height:%.2fpt;" % p.line_spacing.pt
    except Exception: pass
    pre = '<span class="bul">&bull;</span>' if bullet else ""
    return '<p style="text-align:%s;%s">%s%s</p>' % (align, ls, pre, "".join(runs))

def table_html(sh):
    t = sh.table
    out = ['<table style="width:100%;border-collapse:collapse">']
    for ri, row in enumerate(t.rows):
        out.append("<tr>")
        for cell in row.cells:
            fill = shape_fill(cell) or "transparent"
            inner = "".join(para_html(p) for p in cell.text_frame.paragraphs)
            out.append('<td style="background:%s;border-bottom:0.75px solid #D8DDEA;'
                       'padding:5px 9px;vertical-align:middle">%s</td>' % (fill, inner))
        out.append("</tr>")
    out.append("</table>")
    return "".join(out)

def render(path, out_html):
    prs = Presentation(path)
    W, H = I(prs.slide_width), I(prs.slide_height)
    pages = []
    for slide in prs.slides:
        # background
        bg = "#FFFFFF"
        try:
            b = slide.background.fill
            if b.type == 1:
                c = rgb(b.fore_color)
                if c: bg = c
        except Exception: pass
        items = []
        for sh in slide.shapes:
            x, y, w, h = I(sh.left), I(sh.top), I(sh.width), I(sh.height)
            if w <= 0 or h <= 0: continue
            base = "position:absolute;left:%.3fin;top:%.3fin;width:%.3fin;height:%.3fin;" % (x, y, w, h)
            if getattr(sh, "has_table", False) and sh.has_table:
                items.append('<div style="%s">%s</div>' % (base, table_html(sh)))
                continue
            fill = shape_fill(sh)
            st = base
            radius = ""
            try:
                if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                    nm = str(sh.auto_shape_type)
                    if "OVAL" in nm or "ELLIPSE" in nm: radius = "border-radius:50%;"
                    elif "ROUNDED" in nm: radius = "border-radius:0.08in;"
            except Exception: pass
            if fill: st += "background:%s;" % fill
            st += radius
            try:
                ln = sh.line
                if ln.fill.type == 1:
                    c = rgb(ln.color)
                    if c: st += "border:0.75px solid %s;" % c
            except Exception: pass
            inner = ""
            if sh.has_text_frame:
                tf = sh.text_frame
                va = "flex-start"
                try:
                    v = tf.vertical_anchor
                    if v is not None and int(v) == 3: va = "center"
                    elif v is not None and int(v) == 4: va = "flex-end"
                except Exception: pass
                paras = "".join(para_html(p) for p in tf.paragraphs)
                if paras:
                    inner = ('<div style="display:flex;flex-direction:column;justify-content:%s;'
                             'align-items:stretch;height:100%%;width:100%%">%s</div>' % (va, paras))
            items.append('<div style="%s">%s</div>' % (st, inner))
        pages.append('<section style="background:%s">%s</section>' % (bg, "".join(items)))

    doc = """<!doctype html><meta charset="utf-8"><style>
@page { size: %.3fin %.3fin; margin: 0 }
* { box-sizing: border-box }
html,body { margin:0; padding:0; background:#fff }
section { position:relative; width:%.3fin; height:%.3fin; overflow:hidden;
  page-break-after:always; break-after:page; }
section:last-child { page-break-after:auto }
p { margin:0; font-family:Calibri,'DejaVu Sans',Arial,sans-serif; font-size:14pt; color:#1A2340 }
.bul { display:inline-block; width:0.16in; margin-left:-0.16in }
td p { margin:0 }
</style>%s""" % (W, H, W, H, "".join(pages))
    pathlib.Path(out_html).write_text(doc, encoding="utf-8")
    return out_html, W, H

if __name__ == "__main__":
    from playwright.sync_api import sync_playwright
    targets = []
    for src in sys.argv[1:]:
        stem = os.path.splitext(src)[0]
        hp = stem + ".html"
        _, W, Hh = render(src, hp)
        targets.append((os.path.abspath(hp), stem + ".pdf", W, Hh))
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome')
        pg = b.new_page()
        for hp, pdf, W, Hh in targets:
            pg.goto("file://" + hp); pg.wait_for_timeout(400)
            pg.pdf(path=pdf, width="%.3fin" % W, height="%.3fin" % Hh,
                   print_background=True, margin={"top":"0","bottom":"0","left":"0","right":"0"})
            print("wrote", pdf)
        b.close()
