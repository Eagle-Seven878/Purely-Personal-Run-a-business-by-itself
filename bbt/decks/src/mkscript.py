"""Build per-level facilitator scripts as .docx (editable) and .pdf, from the level markdown."""
import re, os, sys, html, pathlib
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO = "/home/user/Purely-Personal-Run-a-business-by-itself/bbt"
INK   = RGBColor(0x14, 0x1C, 0x36)
BRASS = RGBColor(0xA0, 0x71, 0x1E)
DIM   = RGBColor(0x5B, 0x64, 0x80)
WARN  = RGBColor(0xB4, 0x72, 0x0F)

LEVELS = [
 ("BBT-Level-0-Orientation-SCRIPT", "Level 0 · Orientation",
  "Free · 90 minutes · Zoom or in-room", ["level-0-orientation/SCRIPT.md"]),
 ("BBT-Level-1-Masterclass-SCRIPT", "Level 1 · BBT Masterclass",
  "₱1,999 · 240 minutes · live Zoom", ["level-1-masterclass/SCRIPT.md"]),
 ("BBT-Level-2-3-DeepDive-Immersion-SCRIPT", "Levels 2 & 3 · Deep Dive and Immersion",
  "₱15,000 online (Modules 1–6) · ₱75,000 face to face (Modules 1–8)",
  ["level-2-3-deep-dive-immersion/CURRICULUM.md",
   "level-2-3-deep-dive-immersion/MODULE-1-chart-structure.md",
   "level-2-3-deep-dive-immersion/MODULE-2-3-platform-execution.md",
   "level-2-3-deep-dive-immersion/MODULE-4-5-risk-defense.md",
   "level-2-3-deep-dive-immersion/MODULE-6-accumulation.md",
   "level-2-3-deep-dive-immersion/MODULE-7-8-immersion.md"]),
 ("BBT-Level-4-Mentorship-SCRIPT", "Level 4 · Done-With-You Mentorship",
  "USD 3,000 · 12 months", ["level-4-mentorship/SCRIPT.md"]),
 ("BBT-Thrive-Circle-GUIDE", "Continuing · The Thrive Circle",
  "USD 1,999 per year · alumni", ["continuing-thrive-circle/SCRIPT.md"]),
 ("BBT-Train-the-Trainer-SCRIPT", "Train the Trainer",
  "2 days · max 12 candidates · by invitation", ["train-the-trainer/SCRIPT.md"]),
]


def unwrap(md):
    """Join soft-wrapped markdown lines so inline emphasis spanning a line break parses."""
    out, buf, mode = [], None, None
    def flush():
        nonlocal buf, mode
        if buf is not None: out.append(buf)
        buf, mode = None, None
    for raw in md.split("\n"):
        ln = raw.rstrip(); st = ln.strip()
        if st == "":
            flush(); out.append(""); continue
        if st == ">":
            flush(); out.append(">"); continue
        if ln.startswith("> "):
            if mode == "q": buf += " " + st[2:]
            else: flush(); buf = "> " + st[2:]; mode = "q"
            continue
        if st.startswith("#") or st.startswith("|") or st.startswith("---") or st.startswith("```"):
            flush(); out.append(ln); continue
        if re.match(r'^\s*[-*] ', ln) or re.match(r'^\s*\d+\. ', ln):
            flush(); buf = ln; mode = "li"; continue
        if mode in ("p", "li", "q"):
            buf += " " + st
        else:
            flush(); buf = ln; mode = "p"
    flush()
    return "\n".join(out)

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))')
def runs_of(text):
    out = []
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"): out.append((part[2:-2], "b"))
        elif part.startswith("`") and part.endswith("`"):  out.append((part[1:-1], "c"))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: out.append((part[1:-1], "i"))
        elif part.startswith("["):
            m = re.match(r'\[([^\]]+)\]\([^)]+\)', part); out.append((m.group(1) if m else part, ""))
        else: out.append((part, ""))
    return out

def add_runs(par, text, size=10.5, color=None):
    for t, k in runs_of(text):
        r = par.add_run(t.replace("—", "—"))
        r.font.size = Pt(size)
        r.font.name = "Calibri"
        if color: r.font.color.rgb = color
        if k == "b": r.bold = True
        elif k == "i": r.italic = True
        elif k == "c":
            r.font.name = "Consolas"; r.font.size = Pt(size - 0.5)

def build(doc, md, first):
    lines = unwrap(md).split("\n"); i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.startswith("|") and i + 1 < len(lines) and set(lines[i+1].replace("|","").strip()) <= set("-: "):
            hdr = [c.strip() for c in ln.strip("|").split("|")]
            i += 2; rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
            t = doc.add_table(rows=1, cols=len(hdr)); t.style = "Light Grid Accent 1"
            for j, h in enumerate(hdr):
                cell = t.rows[0].cells[j]; cell.text = ""
                add_runs(cell.paragraphs[0], h, 9)
                for r in cell.paragraphs[0].runs: r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for j, c in enumerate(row[:len(hdr)]):
                    cells[j].text = ""; add_runs(cells[j].paragraphs[0], c, 9)
            doc.add_paragraph()
            continue
        if ln.startswith("# "):
            if not first:
                doc.add_page_break()
            p = doc.add_paragraph(); p.space_before = Pt(0)
            add_runs(p, ln[2:], 20); 
            for r in p.runs: r.bold = True; r.font.color.rgb = INK; r.font.name = "Cambria"
            first = False
        elif ln.startswith("## "):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(16)
            add_runs(p, ln[3:], 15)
            for r in p.runs: r.bold = True; r.font.color.rgb = INK; r.font.name = "Cambria"
        elif ln.startswith("### "):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
            add_runs(p, ln[4:], 12)
            for r in p.runs: r.bold = True; r.font.color.rgb = BRASS
        elif ln.startswith("> "):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.32)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, ln[2:], 11.5, INK)
        elif ln.strip() == ">":
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif re.match(r'^\s*[-*] ', ln):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r'^\s*[-*] ', '', ln), 10.5)
        elif re.match(r'^\s*\d+\. ', ln):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r'^\s*\d+\. ', '', ln), 10.5)
        elif ln.startswith("---"):
            pass
        elif ln.strip():
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
            italic = ln.startswith("*") and ln.rstrip().endswith("*") and not ln.startswith("**")
            add_runs(p, ln, 10.5, DIM if italic else None)
        i += 1
    return first

def docx_to_html(name, title, sub, mds):
    """Same content, styled for a print PDF via Chromium."""
    body = []
    for md in mds:
        txt = unwrap(pathlib.Path(os.path.join(REPO, md)).read_text())
        lines = txt.split("\n"); i = 0
        while i < len(lines):
            ln = lines[i].rstrip()
            if ln.startswith("|") and i+1 < len(lines) and set(lines[i+1].replace("|","").strip()) <= set("-: "):
                hdr = [c.strip() for c in ln.strip("|").split("|")]; i += 2; rows=[]
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
                body.append("<table><thead><tr>" + "".join("<th>" + inl(h) for h in hdr) + "</tr></thead><tbody>")
                for r in rows: body.append("<tr>" + "".join("<td>" + inl(c) for c in r[:len(hdr)]) + "</tr>")
                body.append("</tbody></table>"); continue
            if ln.strip() == ">":      i += 1; continue
            if ln.startswith("# "):   body.append("<h1>" + inl(ln[2:]) + "</h1>")
            elif ln.startswith("## "): body.append("<h2>" + inl(ln[3:]) + "</h2>")
            elif ln.startswith("### "):body.append("<h3>" + inl(ln[4:]) + "</h3>")
            elif ln.startswith("> "):  body.append("<blockquote>" + inl(ln[2:]) + "</blockquote>")
            elif re.match(r'^\s*[-*] ', ln): body.append("<li>" + inl(re.sub(r'^\s*[-*] ','',ln)) + "</li>")
            elif ln.startswith("---"): body.append("<hr>")
            elif ln.strip():           body.append("<p>" + inl(ln) + "</p>")
            i += 1
    css = CSS_TEMPLATE
    return (css
            .replace("@@TITLE@@", html.escape(title))
            .replace("@@SUB@@", html.escape(sub))
            .replace("@@BODY@@", "".join(body)))

CSS_TEMPLATE = """<!doctype html><meta charset=utf-8><style>
@page { size: A4; margin: 18mm 16mm }
body{font-family:Calibri,'DejaVu Sans',Arial,sans-serif;font-size:10.5pt;color:#1A2340;line-height:1.45}
h1{font-family:Cambria,'DejaVu Serif',Georgia,serif;font-size:21pt;color:#141C36;margin:0 0 4pt;page-break-before:always}
h1:first-of-type{page-break-before:avoid}
h2{font-family:Cambria,'DejaVu Serif',Georgia,serif;font-size:15pt;color:#141C36;margin:16pt 0 4pt}
h3{font-size:11.5pt;color:#A0711E;margin:12pt 0 3pt}
blockquote{margin:2pt 0 2pt 14pt;color:#141C36;font-size:11pt}
table{border-collapse:collapse;width:100%;margin:8pt 0;font-size:9pt}
th{text-align:left;border-bottom:1px solid #8A92A8;padding:4pt 6pt;color:#5B6480}
td{border-bottom:0.5px solid #D8DDEA;padding:4pt 6pt;vertical-align:top}
li{margin:2pt 0}  hr{border:0;border-top:1px solid #D8DDEA;margin:10pt 0}
code{font-family:Consolas,monospace;font-size:9.5pt;background:#EEF1F8;padding:1pt 3pt}
.cover{page-break-after:always;padding-top:60mm}
.cover .k{font-size:10pt;letter-spacing:2px;color:#A0711E;text-transform:uppercase}
.cover h1{font-size:34pt;page-break-before:avoid;margin:8pt 0}
.cover .s{font-size:13pt;color:#5B6480}
</style>
<div class="cover"><div class="k">Buy &middot; Borrow &middot; Thrive &mdash; facilitator script</div>
<h1>@@TITLE@@</h1><div class="s">@@SUB@@</div></div>@@BODY@@"""

def inl(t):
    t = html.escape(t)
    t = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', t)
    t = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', t)
    t = re.sub(r'`([^`]+?)`', r'<code>\1</code>', t)
    t = re.sub(r'(?<!\*)\*([^*]+?)\*(?!\*)', r'<i>\1</i>', t)
    return t

if __name__ == "__main__":
    htmls = []
    for name, title, sub, mds in LEVELS:
        doc = Document()
        st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
        for s in doc.sections:
            s.left_margin = s.right_margin = Inches(0.85); s.top_margin = s.bottom_margin = Inches(0.8)
        p = doc.add_paragraph(); add_runs(p, "BUY · BORROW · THRIVE — FACILITATOR SCRIPT", 9)
        for r in p.runs: r.font.color.rgb = BRASS; r.bold = True
        p = doc.add_paragraph(); add_runs(p, title, 26)
        for r in p.runs: r.bold = True; r.font.name = "Cambria"; r.font.color.rgb = INK
        p = doc.add_paragraph(); add_runs(p, sub, 12, DIM)
        p = doc.add_paragraph(); add_runs(p, "Verify all platform figures and drawdown percentages before every cohort.", 9.5, WARN)
        doc.add_page_break()
        first = True
        for md in mds:
            first = build(doc, pathlib.Path(os.path.join(REPO, md)).read_text(), first)
        doc.save(name + ".docx"); print("wrote", name + ".docx")
        hp = name + ".html"; pathlib.Path(hp).write_text(docx_to_html(name, title, sub, mds), encoding="utf-8")
        htmls.append((os.path.abspath(hp), name + ".pdf"))

    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = pw.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome')
        pg = b.new_page()
        for hp, pdf in htmls:
            pg.goto("file://" + hp); pg.wait_for_timeout(350)
            pg.pdf(path=pdf, format="A4", print_background=True,
                   margin={"top":"18mm","bottom":"18mm","left":"16mm","right":"16mm"})
            print("wrote", pdf)
        b.close()
